import unittest
import tempfile
from pathlib import Path
from openpyxl import load_workbook

from job_finder import reporter
from tests.fixtures.report_rows import load_fixture


class ReporterTests(unittest.TestCase):
    def _build(self, rows):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        output = reporter.generate_report(rows, Path(directory.name))
        return output, load_workbook(output)

    def test_generate_report_creates_xlsx_with_three_sheets(self):
        output, workbook = self._build(load_fixture())
        self.assertTrue(output.exists())
        self.assertEqual(workbook.sheetnames, ['Pipeline', 'Rejected', 'Run Log'])

    def test_pipeline_layout_and_apply_hyperlink(self):
        rows = load_fixture()
        _, workbook = self._build(rows)
        sheet = workbook['Pipeline']
        headers = [cell.value for cell in sheet[reporter.HEADER_ROW][:len(reporter.PIPELINE_HEADERS)]]
        self.assertEqual(headers, list(reporter.PIPELINE_HEADERS))
        self.assertEqual(sheet.freeze_panes, 'E4')
        self.assertEqual(sheet.cell(row=4, column=1).value, 1)
        self.assertEqual(sheet.cell(row=4, column=2).value, rows[0]['match_score'])

        apply_cell = sheet.cell(row=4, column=reporter.PIPELINE_HEADERS.index('Apply') + 1)
        self.assertIsNotNone(apply_cell.hyperlink)
        self.assertEqual(apply_cell.hyperlink.target, rows[0]['application_url'])
        self.assertLessEqual(len(apply_cell.value), reporter.APPLY_TEXT_LIMIT)

    def test_column_widths_are_not_uniform(self):
        _, workbook = self._build(load_fixture())
        sheet = workbook['Pipeline']
        self.assertEqual(sheet.column_dimensions['A'].width, 5)
        self.assertEqual(sheet.column_dimensions['D'].width, 24)
        self.assertEqual(sheet.column_dimensions['E'].width, 38)
        self.assertEqual(sheet.column_dimensions['P'].width, 46)
        self.assertEqual(sheet.column_dimensions['R'].width, 46)
        self.assertTrue(sheet.cell(row=4, column=5).alignment.wrap_text)

    def test_rejected_rows_are_separated_from_pipeline(self):
        rows = load_fixture() + [dict(load_fixture()[0], job_id='sample-2', match_score=12,
                                      recommendation='DO NOT INCLUDE')]
        _, workbook = self._build(rows)
        self.assertEqual(workbook['Pipeline'].cell(row=4, column=4).value, rows[0]['company'])
        self.assertIsNone(workbook['Pipeline'].cell(row=5, column=1).value)
        self.assertEqual(workbook['Rejected'].cell(row=4, column=1).value, 12)

    def test_run_log_reports_verification_rates(self):
        rows = load_fixture() + [dict(load_fixture()[0], job_id='sample-2',
                                      recommendation='DO NOT INCLUDE', salary='Not disclosed',
                                      posting_date_verified=False, location_verified=False,
                                      cloud_devops_match='Low')]
        rows[0].update(posting_date_verified=True, location_verified=True, salary='18 LPA')
        run = {'run_boards_fetched': 2, 'run_listings_fetched': 5, 'run_india_relevant': 3,
               'run_new_rows': 1, 'run_updated_rows': 1, 'run_rows_in_pipeline': 1,
               'run_rejection_reasons': {'below_threshold': 1}, 'score_histogram': {'60-69': 1}}
        metrics = dict(reporter.run_log_metrics(rows, run))
        self.assertEqual(metrics['db_rows_total'], 2)
        self.assertEqual(metrics['db_rows_in_pipeline'], 1)
        self.assertEqual(metrics['db_rows_rejected'], 1)
        self.assertEqual(metrics['db_posting_date_verified_rate'], 0.5)
        self.assertEqual(metrics['db_location_verified_rate'], 0.5)
        self.assertEqual(metrics['db_salary_parsed_rate'], 0.5)
        self.assertEqual(metrics['db_cloud_fit_distribution'], 'High: 1, Low: 1')
        self.assertEqual(metrics['run_rejection_reasons'], {'below_threshold': 1})
        self.assertEqual(metrics['score_histogram'], {'60-69': 1})

        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        workbook = load_workbook(reporter.generate_report(rows, Path(directory.name), run))
        sheet = workbook['Run Log']
        written = {sheet.cell(row=row, column=1).value: sheet.cell(row=row, column=2).value
                   for row in range(4, 4 + len(metrics))}
        self.assertEqual(written['db_rows_total'], 2)
        self.assertEqual(written['db_salary_parsed_rate'], 0.5)
        self.assertEqual(written['score_histogram'], '{"60-69": 1}')

    def test_rejection_reasons_counts_excessive_experience(self):
        rows = [{
            'company': 'Acme',
            'title': 'Software Engineer',
            'location': 'Hyderabad, India',
            'experience': '6 yrs',
            'recommendation': 'DO NOT INCLUDE',
            'seniority': 'MID',
            'location_verified': True,
            'freshness': '1-3 days',
            'source': 'Lever board API (acme)',
        }]
        reasons = reporter.rejection_reasons(rows)
        self.assertEqual(reasons['too_much_experience'], 1)
        self.assertEqual(sum(reasons.values()), 1)


if __name__ == '__main__':
    unittest.main()
