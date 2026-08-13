from __future__ import annotations

import json
import re
from collections import Counter
from datetime import date
from pathlib import Path
from typing import Any, Sequence

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from .models import Job
from .scoring import classify_seniority, is_excluded, is_stale, location_points, salary_lpa

HEADER_ROW = 3
DATA_ROW = 4

TITLE_FILL = PatternFill('solid', start_color='0F2A43', end_color='0F2A43')
HEADER_FILL = PatternFill('solid', start_color='146C94', end_color='146C94')
AMBER_FILL = PatternFill('solid', start_color='FDE68A', end_color='FDE68A')

STATUS_CHOICES = ('Not Applied', 'Applied', 'Referral Requested', 'Recruiter Replied',
                  'Screen Scheduled', 'Round 1', 'Round 2', 'Offer', 'Rejected', 'Ghosted')

PIPELINE_HEADERS = ('#', 'Score', 'Action', 'Company', 'Role', 'Experience', 'Location',
                    'Mode', 'Salary', 'Salary (LPA)', 'Posted', 'Age', 'Matching Skills', 'Skill Gaps',
                    'Cloud Fit', 'Why', 'Concerns', 'Apply', 'Status', 'Applied On', 'Notes')

# Anything unlisted falls back to DEFAULT_WIDTH; the old sheet set all 30 to 15,
# which is why long text columns rendered on top of each other.
DEFAULT_WIDTH = 14
PIPELINE_WIDTHS = {'#': 5, 'Score': 7, 'Action': 20, 'Company': 24, 'Role': 38,
                   'Experience': 13, 'Location': 22, 'Mode': 10, 'Salary': 18, 'Salary (LPA)': 12,
                   'Posted': 12, 'Age': 14, 'Matching Skills': 30, 'Skill Gaps': 24,
                   'Cloud Fit': 10, 'Why': 46, 'Concerns': 32, 'Apply': 46,
                   'Status': 20, 'Applied On': 13, 'Notes': 30}
PIPELINE_WRAP = ('Role', 'Matching Skills', 'Skill Gaps', 'Why', 'Concerns', 'Notes')

REJECTED_HEADERS = ('Score', 'Company', 'Role', 'Location', 'Experience', 'Posted',
                    'Cloud Fit', 'Why Rejected', 'Eligibility', 'Confidence', 'Source', 'Apply')
REJECTED_WIDTHS = {'Score': 7, 'Company': 24, 'Role': 38, 'Location': 22, 'Experience': 13,
                   'Posted': 12, 'Cloud Fit': 10, 'Why Rejected': 46, 'Eligibility': 22,
                   'Confidence': 12, 'Source': 26, 'Apply': 46}
REJECTED_WRAP = ('Role', 'Why Rejected')

APPLY_TEXT_LIMIT = 45


def _score(row: dict) -> int:
    value = row.get('match_score') or 0
    return int(value) if isinstance(value, (int, float)) else 0


def _is_rejected(row: dict) -> bool:
    return row.get('recommendation') == 'DO NOT INCLUDE'


def _text(value: Any) -> str:
    if value is None:
        return ''
    if isinstance(value, (list, tuple)):
        return ', '.join(str(item) for item in value)
    return str(value)


def _truthy(value: Any) -> bool:
    if isinstance(value, str):
        return value.strip().lower() in {'1', 'true', 'yes', 't'}
    return bool(value)


def _write_frame(sheet: Worksheet, title: str, headers: Sequence[str], widths: dict,
                 wrap: Sequence[str], row_count: int) -> None:
    last_column = get_column_letter(len(headers))
    sheet.merge_cells(f'A1:{last_column}1')
    heading = sheet['A1']
    heading.value = f'{title} — Aryan Kumar Job Finder — {date.today().isoformat()}'
    heading.fill = TITLE_FILL
    heading.font = Font(bold=True, color='FFFFFF', size=14)
    heading.alignment = Alignment(horizontal='left', vertical='center')
    sheet.row_dimensions[1].height = 24

    for index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=HEADER_ROW, column=index, value=header)
        cell.fill = HEADER_FILL
        cell.font = Font(bold=True, color='FFFFFF')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = widths.get(header, DEFAULT_WIDTH)

    last_row = HEADER_ROW + max(row_count, 1)
    sheet.auto_filter.ref = f'A{HEADER_ROW}:{last_column}{last_row}'

    wrap_columns = {get_column_letter(index) for index, header in enumerate(headers, start=1) if header in wrap}
    for row in sheet.iter_rows(min_row=DATA_ROW, max_row=HEADER_ROW + max(row_count, 0), max_col=len(headers)):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=cell.column_letter in wrap_columns)


def _apply_cell(sheet: Worksheet, row_index: int, column_index: int, url: str) -> None:
    cell = sheet.cell(row=row_index, column=column_index)
    if not url:
        return
    cell.value = url if len(url) <= APPLY_TEXT_LIMIT else url[:APPLY_TEXT_LIMIT - 1] + '…'
    cell.hyperlink = url
    cell.style = 'Hyperlink'


def _pipeline_sheet(workbook: Workbook, rows: list[dict]) -> None:
    sheet = workbook.create_sheet('Pipeline')
    ordered = sorted(rows, key=lambda row: -_score(row))

    for offset, row in enumerate(ordered):
        excel_row = DATA_ROW + offset
        values = (offset + 1, _score(row), _text(row.get('recommendation')), _text(row.get('company')),
                  _text(row.get('title')), _text(row.get('experience')), _text(row.get('location')),
                  _text(row.get('work_mode')), _text(row.get('salary')),
                  salary_lpa(_text(row.get('salary'))), _text(row.get('posting_date')),
                  _text(row.get('freshness')),
                  _text(row.get('matching_skills')),
                  _text(row.get('skill_gaps')), _text(row.get('cloud_devops_match')),
                  _text(row.get('why_it_matches')), _text(row.get('potential_concerns')),
                  None, 'Not Applied', None, None)
        for index, value in enumerate(values, start=1):
            sheet.cell(row=excel_row, column=index, value=value)
        _apply_cell(sheet, excel_row, PIPELINE_HEADERS.index('Apply') + 1,
                    _text(row.get('application_url') or row.get('original_url')))

    _write_frame(sheet, 'Pipeline', PIPELINE_HEADERS, PIPELINE_WIDTHS, PIPELINE_WRAP, len(ordered))
    sheet.freeze_panes = 'E4'

    last_row = HEADER_ROW + max(len(ordered), 1)
    score_range = f'B{DATA_ROW}:B{last_row}'
    sheet.conditional_formatting.add(score_range, ColorScaleRule(
        start_type='num', start_value=45, start_color='F87171',
        mid_type='num', mid_value=68, mid_color='FDE68A',
        end_type='num', end_value=88, end_color='4ADE80'))

    experience_column = get_column_letter(PIPELINE_HEADERS.index('Experience') + 1)
    sheet.conditional_formatting.add(
        f'{experience_column}{DATA_ROW}:{experience_column}{last_row}',
        CellIsRule(operator='equal', formula=['"Unknown"'], fill=AMBER_FILL))

    status_column = get_column_letter(PIPELINE_HEADERS.index('Status') + 1)
    validation = DataValidation(type='list', formula1='"' + ','.join(STATUS_CHOICES) + '"',
                                allow_blank=True, showDropDown=False)
    validation.error = 'Pick a status from the list.'
    validation.prompt = 'Track where this application stands.'
    sheet.add_data_validation(validation)
    validation.add(f'{status_column}{DATA_ROW}:{status_column}{last_row}')


def _rejected_sheet(workbook: Workbook, rows: list[dict]) -> None:
    sheet = workbook.create_sheet('Rejected')
    ordered = sorted(rows, key=lambda row: -_score(row))

    for offset, row in enumerate(ordered):
        excel_row = DATA_ROW + offset
        values = (_score(row), _text(row.get('company')), _text(row.get('title')),
                  _text(row.get('location')), _text(row.get('experience')),
                  _text(row.get('posting_date')), _text(row.get('cloud_devops_match')),
                  _text(row.get('potential_concerns')), _text(row.get('eligibility_status')),
                  _text(row.get('confidence')), _text(row.get('source')), None)
        for index, value in enumerate(values, start=1):
            sheet.cell(row=excel_row, column=index, value=value)
        _apply_cell(sheet, excel_row, REJECTED_HEADERS.index('Apply') + 1,
                    _text(row.get('application_url') or row.get('original_url')))

    _write_frame(sheet, 'Rejected', REJECTED_HEADERS, REJECTED_WIDTHS, REJECTED_WRAP, len(ordered))
    sheet.freeze_panes = 'D4'


def _rate(count: int, total: int) -> float:
    return round(count / total, 2) if total else 0.0


def rejection_reasons(rows: list[dict]) -> dict[str, int]:
    """Classify rejected rows by the same gate order used by scoring.score()."""
    reasons = {key: 0 for key in ('excluded_role', 'too_senior', 'outside_location', 'stale', 'below_threshold')}
    for row in rows:
        if not _is_rejected(row):
            continue
        job = Job(
            company=_text(row.get('company')),
            title=_text(row.get('title')),
            location=_text(row.get('location')) or 'Unknown',
            source=_text(row.get('source')),
        )
        job.freshness = _text(row.get('freshness'))
        if is_excluded(job):
            reasons['excluded_role'] += 1
        elif _text(row.get('seniority')) in {'STAFF', 'PRINCIPAL', 'LEAD', 'MANAGER', 'DIRECTOR'} or classify_seniority(job.title) in {'STAFF', 'PRINCIPAL', 'LEAD', 'MANAGER', 'DIRECTOR'}:
            reasons['too_senior'] += 1
        elif _truthy(row.get('location_verified')) and location_points(job) == 0:
            reasons['outside_location'] += 1
        elif is_stale(job):
            reasons['stale'] += 1
        else:
            reasons['below_threshold'] += 1
    return reasons


def run_log_metrics(rows: list[dict], run_metrics: dict | None = None) -> list[tuple[str, Any]]:
    """Whole-database quality metrics plus the separately persisted current run."""
    total = len(rows)
    pipeline = [row for row in rows if not _is_rejected(row)]
    salary_parsed = sum(1 for row in rows
                        if _text(row.get('salary')).strip() not in {'', 'Not disclosed', 'Unknown'})
    distribution = Counter(_text(row.get('cloud_devops_match')) or 'Unknown' for row in rows)
    metrics = [
        ('db_rows_total', total),
        ('db_rows_in_pipeline', len(pipeline)),
        ('db_rows_rejected', total - len(pipeline)),
        ('db_posting_date_verified_rate', _rate(sum(1 for row in rows if _truthy(row.get('posting_date_verified'))), total)),
        ('db_location_verified_rate', _rate(sum(1 for row in rows if _truthy(row.get('location_verified'))), total)),
        ('db_salary_parsed_rate', _rate(salary_parsed, total)),
        ('db_cloud_fit_distribution', ', '.join(f'{label}: {count}' for label, count in sorted(distribution.items())) or 'none'),
        ('db_rejection_reasons', rejection_reasons(rows)),
    ]
    metrics.extend((key, (run_metrics or {}).get(key, {} if key in {'run_rejection_reasons', 'score_histogram'} else 0))
                   for key in ('run_boards_fetched', 'run_listings_fetched', 'run_india_relevant',
                               'run_new_rows', 'run_updated_rows', 'run_rows_in_pipeline',
                               'run_rejection_reasons', 'score_histogram'))
    return metrics


TARGETS = {'db_posting_date_verified_rate': '> 0.90', 'db_salary_parsed_rate': '> 0.20',
           'db_cloud_fit_distribution': 'want a spread, not one bucket'}


def _run_log_sheet(workbook: Workbook, rows: list[dict], run_metrics: dict | None = None) -> None:
    sheet = workbook.create_sheet('Run Log')
    headers = ('Metric', 'Value', 'Target')
    widths = {'Metric': 32, 'Value': 46, 'Target': 30}
    metrics = run_log_metrics(rows, run_metrics)
    for offset, (name, value) in enumerate(metrics):
        excel_row = DATA_ROW + offset
        sheet.cell(row=excel_row, column=1, value=name).font = Font(bold=True)
        sheet.cell(row=excel_row, column=2, value=json.dumps(value, sort_keys=True) if isinstance(value, dict) else value)
        sheet.cell(row=excel_row, column=3, value=TARGETS.get(name, ''))
    _write_frame(sheet, 'Run Log', headers, widths, ('Value', 'Target'), len(metrics))
    sheet.freeze_panes = 'A4'


def generate_report(rows: list[dict], report_dir: Path, run_metrics: dict | None = None) -> Path:
    report_dir.mkdir(parents=True, exist_ok=True)
    output = report_dir / f'Aryan_Job_Report_{date.today().isoformat()}.xlsx'

    rows = [row for row in rows if isinstance(row, dict)]
    workbook = Workbook()
    workbook.remove(workbook.active)
    _pipeline_sheet(workbook, [row for row in rows if not _is_rejected(row)])
    _rejected_sheet(workbook, [row for row in rows if _is_rejected(row)])
    _run_log_sheet(workbook, rows, run_metrics)
    workbook.save(output)
    return output
